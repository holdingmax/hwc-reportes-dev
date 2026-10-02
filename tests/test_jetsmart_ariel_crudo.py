"""Lectura del archivo CRUDO de Ariel (su rendicion mensual, hoja "Venta")
en vez de la hoja ARIEL del archivo de trabajo de Anita.

Correr con:
    python -m unittest tests.test_jetsmart_ariel_crudo -v

Los tests de formato (nombre de hoja, filas en blanco arriba, errores) usan
archivos chicos armados en el momento y no necesitan fixtures. Los de
equivalencia arman una rendicion cruda con los datos reales de la hoja
ARIEL de julio 2026 (tests/fixtures/jetsmart/, no versionado) y verifican
que leerla dé exactamente lo mismo que la hoja ARIEL armada a mano, y que
el cruce completo no cambie.
"""

import io
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import Workbook

from src.jetsmart_builder import detect_jetsmart_period, load_jetsmart_export
from src.config import (
    COL_JS_CLIENTE,
    COL_JS_CREACION,
    COL_JS_DESTINO,
    COL_JS_EMPRESA,
    COL_JS_ESTADO,
    COL_JS_GUIA,
    COL_JS_KGS,
    COL_JS_ORIGEN,
    COL_JS_PRIORIDAD,
    JETSMART_STATIONS,
)
from src.jetsmart_cruce import (
    _AR_COLUMNS,
    EXC_DOBLE,
    EXC_DUPLICADA,
    INC_PENDIENTE_ANTERIOR,
    cruzar_jetsmart,
    has_ariel_raw_file,
    has_ariel_sheet,
    load_ariel,
    load_ariel_raw,
)

FIX = Path(__file__).resolve().parent / "fixtures" / "jetsmart"
TRABAJO_JUN, TRABAJO_JUL = FIX / "LIQUIDACION_06-2026.xlsx", FIX / "LIQUIDACION_07-2026.xlsx"
# Rendiciones reales que mando Ariel (crudas, sin tocar).
RENDICION_AGO = FIX / "Ariel_Rendicion_agosto_2026.xlsx"
RENDICION_SEP = FIX / "Ariel_Rendicion_septiembre_2026.xlsx"

FILAS_CHICAS = [
    [827, 50200474, 94542, datetime(2026, 7, 1), 3039, "AMDM S.R.L", "JULIO BURGOA", "AEP", "BRC", 1, 8, 8, 736, 10000, "ART. CONSUMO"],
    [827, 50200533, 94549, datetime(2026, 7, 1), 3039, "ENVIO FACIL SRL", "ENVIO FACIL", "aep ", "BRC", 1, 12, 12, 633, 12000.5, "DOC"],
    [827, 50201502, 94562, "ANULADA", None, None, None, None, None, None, None, None, None, None, None],
    [827, 50202891, 94562, datetime(2026, 7, 3), 3031, "OCA LOG S.A", "OCA LOG", "AEP", "CRD", 8, 145, 150, 240, 36000, "DOC"],
    [827, 30012511, "ANULADA", None, None, None, None, None, None, None, None, None, None, None, None],
]


def _rendicion(filas, *, hoja="Venta", blancos=2, encabezado=_AR_COLUMNS, otras_hojas=("Mane EZE", "mane MDZ"),
               col_inicio=1, con_total=True, manifiesto_primero=False) -> io.BytesIO:
    """Arma en memoria un archivo con la forma de la rendicion de Ariel.
    manifiesto_primero: agrega antes de "Venta" una hoja con el mismo
    encabezado y otras guias, para probar que solo se lee "Venta"."""
    wb = Workbook()
    ws = wb.active
    if manifiesto_primero:
        ws.title = "Mane EZE copia"
        ws.append(list(_AR_COLUMNS))
        ws.append([827, 1, 11111, datetime(2026, 7, 1)] + [None] * 11)
        ws = wb.create_sheet(hoja)
    else:
        ws.title = hoja
    fila = 1 + blancos
    for j, h in enumerate(encabezado):
        ws.cell(row=fila, column=col_inicio + j, value=h)
    for i, valores in enumerate(filas, start=1):
        for j, v in enumerate(valores):
            ws.cell(row=fila + i, column=col_inicio + j, value=v)
    if con_total:  # pie de totales, como trae el original
        ws.cell(row=fila + len(filas) + 2, column=col_inicio + _AR_COLUMNS.index("Kg"), value=999)
    for nombre in otras_hojas:  # manifiestos: no se tocan
        m = wb.create_sheet(nombre)
        m.append(["Fecha", "Vuelo", "Mani", "Guias"])
        m.append([datetime(2026, 7, 1), "WJ3810", "33199-Y", 0])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestFormatoRendicion(unittest.TestCase):
    def test_hoja_venta_sin_importar_mayusculas(self):
        for nombre in ("Venta", "VENTA", "venta", " Venta "):
            with self.subTest(hoja=nombre):
                df, anuladas = load_ariel_raw(_rendicion(FILAS_CHICAS, hoja=nombre))
                self.assertEqual(list(df["Nro de guía"]), [94542, 94549, 94562])
                self.assertEqual(anuladas, 2)

    def test_detecta_el_encabezado_con_0_1_y_2_filas_en_blanco(self):
        for blancos in (0, 1, 2):
            with self.subTest(filas_en_blanco=blancos):
                df, _ = load_ariel_raw(_rendicion(FILAS_CHICAS, blancos=blancos))
                self.assertEqual(len(df), 3)
                self.assertEqual(list(df.columns), _AR_COLUMNS)

    def test_tabla_que_arranca_en_otra_columna(self):
        df, _ = load_ariel_raw(_rendicion(FILAS_CHICAS, col_inicio=2))
        self.assertEqual(list(df.columns), _AR_COLUMNS)
        self.assertEqual(len(df), 3)

    def test_encabezado_con_tildes_o_espacios_distintos(self):
        variante = [h.upper() for h in _AR_COLUMNS]
        variante[_AR_COLUMNS.index("Nro de guía")] = " Nro de Guia "
        df, _ = load_ariel_raw(_rendicion(FILAS_CHICAS, encabezado=variante))
        self.assertEqual(list(df.columns), _AR_COLUMNS)

    def test_misma_limpieza_que_la_hoja_ariel(self):
        df, _ = load_ariel_raw(_rendicion(FILAS_CHICAS))
        # anuladas fuera (incluida la que reusa el numero 94562), origen
        # normalizado, fecha como fecha, sin el pie de totales
        self.assertEqual(df.loc[df["Nro de guía"] == 94562, "AWB"].tolist(), [50202891])
        self.assertEqual(df["Origen"].tolist(), ["AEP", "AEP", "AEP"])
        self.assertTrue(pd.api.types.is_datetime64_any_dtype(df["Fecha de guia"]))

    def test_las_otras_hojas_no_se_leen(self):
        # otra hoja (antes de "Venta") con el mismo encabezado no se mezcla
        df, _ = load_ariel_raw(_rendicion(FILAS_CHICAS, manifiesto_primero=True))
        self.assertEqual(list(df["Nro de guía"]), [94542, 94549, 94562])

    def test_error_sin_hoja_venta(self):
        archivo = _rendicion(FILAS_CHICAS, hoja="Rendicion")
        self.assertFalse(has_ariel_raw_file(archivo))
        with self.assertRaisesRegex(ValueError, "no tiene la hoja Venta"):
            load_ariel_raw(archivo)

    def test_error_sin_encabezado(self):
        archivo = _rendicion(FILAS_CHICAS, encabezado=["Col A", "Col B", "Col C"])
        self.assertFalse(has_ariel_raw_file(archivo))
        with self.assertRaisesRegex(ValueError, "no tiene el encabezado esperado"):
            load_ariel_raw(archivo)

    def test_error_encabezado_demasiado_abajo(self):
        archivo = _rendicion(FILAS_CHICAS, blancos=25)
        with self.assertRaisesRegex(ValueError, "no tiene el encabezado esperado"):
            load_ariel_raw(archivo)

    def test_rendicion_valida_no_es_archivo_de_trabajo(self):
        archivo = _rendicion(FILAS_CHICAS)
        self.assertTrue(has_ariel_raw_file(archivo))
        self.assertFalse(has_ariel_sheet(archivo))


@unittest.skipUnless(TRABAJO_JUN.exists() and TRABAJO_JUL.exists(), "faltan los archivos de tests/fixtures/jetsmart/")
class TestEquivalenciaConHojaAriel(unittest.TestCase):
    """La rendicion cruda armada con los datos reales de julio tiene que dar
    lo mismo que la hoja ARIEL que Anita armo a mano con esos datos."""

    @classmethod
    def setUpClass(cls):
        with pd.ExcelFile(TRABAJO_JUL) as xls:
            hoja = next(s for s in xls.sheet_names if s.strip() == "ARIEL")
            original = pd.read_excel(xls, sheet_name=hoja)[_AR_COLUMNS]
        filas = original.astype(object).where(original.notna(), None).values.tolist()
        cls.crudo = _rendicion(filas, blancos=2, con_total=False)
        cls.de_hoja, cls.anuladas_hoja = load_ariel(TRABAJO_JUL)
        cls.de_crudo, cls.anuladas_crudo = load_ariel_raw(cls.crudo)

    def test_mismas_guias_y_anuladas(self):
        self.assertEqual(self.anuladas_crudo, self.anuladas_hoja)
        self.assertEqual(list(self.de_crudo["Nro de guía"]), list(self.de_hoja["Nro de guía"]))

    def test_mismos_valores(self):
        cols = ["Nro de guía", "Origen", "Destino", "Cliente"]
        pd.testing.assert_frame_equal(self.de_crudo[cols], self.de_hoja[cols])
        for col in ("Kg", "Ingreso total flete"):
            pd.testing.assert_series_equal(
                pd.to_numeric(self.de_crudo[col]), pd.to_numeric(self.de_hoja[col]), check_dtype=False,
            )
        pd.testing.assert_series_equal(self.de_crudo["Fecha de guia"], self.de_hoja["Fecha de guia"])

    def test_cruce_identico(self):
        export_jul, export_jun = load_jetsmart_export(TRABAJO_JUL), load_jetsmart_export(TRABAJO_JUN)
        ariel_jun, _ = load_ariel(TRABAJO_JUN)
        periodo, _ = detect_jetsmart_period(export_jul)
        con_hoja = cruzar_jetsmart(export_jul, self.de_hoja, periodo, export_jun, ariel_jun)
        con_crudo = cruzar_jetsmart(export_jul, self.de_crudo, periodo, export_jun, ariel_jun)
        self.assertEqual(con_crudo.resumen, con_hoja.resumen)
        pd.testing.assert_frame_equal(con_crudo.excepciones, con_hoja.excepciones)
        self.assertEqual(set(con_crudo.incluidas["# Guía"]), set(con_hoja.incluidas["# Guía"]))
        pd.testing.assert_series_equal(
            pd.to_numeric(con_crudo.incluidas["$ Prioridad"]), pd.to_numeric(con_hoja.incluidas["$ Prioridad"]),
            check_dtype=False,
        )



def _lectura_a_mano(path) -> dict:
    """Lectura independiente de la hoja de ventas, celda por celda con
    openpyxl (sin pasar por load_ariel_raw), para comparar contra ella."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    fila_hdr = next(i for i, r in enumerate(filas) if r and r[0] == "Prefihjo")
    validas = [
        r for r in filas[fila_hdr + 1:]
        if isinstance(r[2], (int, float))
        and not any(str(c).strip().upper() == "ANULADA" for c in r if c is not None)
    ]
    anuladas = sum(
        any(str(c).strip().upper() == "ANULADA" for c in r if c is not None) for r in filas[fila_hdr + 1:]
    )
    return {
        "hoja": ws.title, "filas_en_blanco": fila_hdr, "guias": [int(r[2]) for r in validas],
        "anuladas": anuladas, "kg": sum(r[10] for r in validas), "ingreso": sum(r[13] for r in validas),
        "max_columna": ws.max_column,
    }


def _export_vacio() -> pd.DataFrame:
    # Sin el export de guias de agosto/septiembre (no lo tenemos): solo se
    # validan las reglas del cruce que dependen de los archivos de Ariel.
    return pd.DataFrame(columns=[
        COL_JS_CREACION, COL_JS_GUIA, COL_JS_CLIENTE, COL_JS_ORIGEN, COL_JS_DESTINO,
        COL_JS_ESTADO, COL_JS_KGS, COL_JS_PRIORIDAD, COL_JS_EMPRESA,
    ])


@unittest.skipUnless(RENDICION_AGO.exists() and RENDICION_SEP.exists(), "faltan las rendiciones reales de Ariel")
class TestRendicionesReales(unittest.TestCase):
    """Rendiciones reales de agosto y septiembre 2026, tal como las mando
    Ariel. Difieren justo en lo que hay que tolerar: nombre de la hoja
    ("VENTA" vs "Venta") y filas en blanco arriba del encabezado (1 vs 2)."""

    @classmethod
    def setUpClass(cls):
        cls.mano = {"ago": _lectura_a_mano(RENDICION_AGO), "sep": _lectura_a_mano(RENDICION_SEP)}
        cls.ago, cls.anuladas_ago = load_ariel_raw(RENDICION_AGO)
        cls.sep, cls.anuladas_sep = load_ariel_raw(RENDICION_SEP)

    def test_los_dos_archivos_difieren_en_hoja_y_offset(self):
        self.assertEqual((self.mano["ago"]["hoja"], self.mano["ago"]["filas_en_blanco"]), ("VENTA", 1))
        self.assertEqual((self.mano["sep"]["hoja"], self.mano["sep"]["filas_en_blanco"]), ("Venta", 2))

    def test_detecta_la_hoja_venta(self):
        for path in (RENDICION_AGO, RENDICION_SEP):
            with self.subTest(archivo=path.name):
                self.assertTrue(has_ariel_raw_file(path))
                self.assertFalse(has_ariel_sheet(path))

    def test_columnas(self):
        # septiembre trae 26 columnas en la hoja (formato vacio): solo
        # quedan las 15 de la tabla.
        self.assertEqual(self.mano["sep"]["max_columna"], 26)
        for df in (self.ago, self.sep):
            self.assertEqual(list(df.columns), _AR_COLUMNS)

    def test_datos_iguales_a_la_lectura_a_mano(self):
        for mes, df, anuladas in (("ago", self.ago, self.anuladas_ago), ("sep", self.sep, self.anuladas_sep)):
            mano = self.mano[mes]
            with self.subTest(mes=mes):
                self.assertEqual(list(df["Nro de guía"]), mano["guias"])
                self.assertEqual(anuladas, mano["anuladas"])
                self.assertAlmostEqual(float(pd.to_numeric(df["Kg"]).sum()), mano["kg"], places=6)
                self.assertAlmostEqual(float(pd.to_numeric(df["Ingreso total flete"]).sum()), mano["ingreso"], places=6)
        self.assertEqual((len(self.ago), self.anuladas_ago), (1561, 1))
        self.assertEqual((len(self.sep), self.anuladas_sep), (1978, 5))
        self.assertAlmostEqual(float(pd.to_numeric(self.ago["Ingreso total flete"]).sum()), 90421072.15, places=2)
        self.assertAlmostEqual(float(pd.to_numeric(self.sep["Ingreso total flete"]).sum()), 105435776.57, places=2)

    def test_valores_con_sentido(self):
        for mes, df, inicio, fin in (("ago", self.ago, "2026-08-01", "2026-09-01"), ("sep", self.sep, "2026-08-31", "2026-09-30")):
            with self.subTest(mes=mes):
                self.assertFalse(df["Fecha de guia"].isna().any())
                self.assertGreaterEqual(df["Fecha de guia"].min(), pd.Timestamp(inicio))
                self.assertLessEqual(df["Fecha de guia"].max(), pd.Timestamp(fin))
                self.assertTrue(set(df["Origen"]) <= set(JETSMART_STATIONS))
                for col in ("Kg", "Ingreso total flete"):
                    self.assertFalse(pd.to_numeric(df[col], errors="coerce").isna().any())

    @unittest.skipUnless(TRABAJO_JUL.exists(), "falta el archivo de trabajo de julio")
    def test_cruce_agosto_contra_julio(self):
        # Sin el export de agosto: todo Ariel queda "solo en Ariel". Lo que
        # se valida es lo que depende de los archivos de Ariel/julio.
        ariel_jul, _ = load_ariel(TRABAJO_JUL)
        r = cruzar_jetsmart(_export_vacio(), self.ago, ("AGO", "26"), load_jetsmart_export(TRABAJO_JUL), ariel_jul)
        exc = r.excepciones
        # 96640, la pendiente de declarar de julio, Ariel la declara en agosto
        self.assertEqual(r.resumen["motivos_inclusion"].get(INC_PENDIENTE_ANTERIOR), 1)
        self.assertIn(96640, set(r.incluidas["# Guía"]))
        # 97479 viene dos veces en la rendicion (misma fila): revision manual
        self.assertEqual(set(exc.loc[exc["Tipo"] == EXC_DUPLICADA, "Nro guía"]), {97479})
        # nada de agosto estaba ya declarado en julio
        self.assertFalse((exc["Tipo"] == EXC_DOBLE).any())

    def test_cruce_septiembre_contra_agosto(self):
        r = cruzar_jetsmart(_export_vacio(), self.sep, ("SEP", "26"), None, self.ago)
        exc = r.excepciones
        # 4 guias que Ariel declaro en agosto Y en septiembre: el cruce las
        # frena (no entran) para no liquidarlas dos veces.
        dobles = set(exc.loc[exc["Tipo"] == EXC_DOBLE, "Nro guía"])
        self.assertEqual(dobles, {98172, 98177, 98182, 98188})
        self.assertFalse(dobles & set(r.incluidas["# Guía"]))
        self.assertFalse((exc["Tipo"] == EXC_DUPLICADA).any())


if __name__ == "__main__":
    unittest.main()
