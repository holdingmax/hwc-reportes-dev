"""Armado de la liquidacion de JetSmart.

JetSmart no encaja en ninguno de los dos flujos existentes:
- Parte de otro export (el del sistema de guias de Handyway filtrado por
  Empresa, una fila por guia con KGs y $ Prioridad), no de
  "archivo original.xlsx".
- Su liquidacion (hoja LIQUIDACION de la planilla de Anita) es Ventas
  totales / IVA / Ventas Netas / Comisiones / GHA Services / IIBB / Total a
  entregar a WCS, mas la hoja CVLP. No comparte formula con el Resumen
  Facturacion de LATAM (sub-facturas LA / 4M), asi que no reusa
  build_latam_resumen: generalizarla para cubrir los dos casos la
  convertiria en dos funciones distintas pegadas por un if.

Formulas relevadas celda por celda de LIQ_ECS_07-2026_00000005.xlsx y
validadas al centavo en tests/test_jetsmart_regresion.py.

QUE guias entran lo define el cruce contra el archivo de Ariel
(jetsmart_cruce.py): app.py y main.py le pasan a build_jetsmart_guias las
guias ya cruzadas, con los valores de Ariel. Las funciones de este modulo no
saben nada del cruce: arman la liquidacion sobre cualquier lista de guias en
formato export. Fuera de alcance: la validacion contra AFIP.
"""

import calendar

import pandas as pd
from openpyxl.styles import Font

from src.config import (
    COL_JS_CLIENTE,
    COL_JS_CREACION,
    COL_JS_DESTINO,
    COL_JS_EMPRESA,
    COL_JS_ESTADO,
    COL_JS_GUIA,
    COL_JS_KGS,
    COL_JS_ORIGEN,
    COL_JS_PRIORIDAD,
    JETSMART_COMISION_DOMESTICA_RATE,
    JETSMART_COMISION_INTER_USD_POR_VUELO,
    JETSMART_IIBB,
    JETSMART_IVA_RATE,
    JETSMART_IVA_SERVICIOS,
    JETSMART_STATIONS,
    JETSMART_TARIFA_CONFIRMADA_STATIONS,
    JETSMART_TARIFA_USD_KG_DEFAULT,
    JETSMART_TARIFA_USD_KG_POR_ESTACION,
)
from src.period_utils import Period

# Mismas abreviaturas que usa period_utils/db.py para el resto de las
# aerolineas, asi el periodo de JetSmart se guarda y se filtra igual.
_MES_ABREV = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]

_REQUIRED_COLUMNS = [
    COL_JS_CREACION, COL_JS_GUIA, COL_JS_CLIENTE, COL_JS_ORIGEN, COL_JS_DESTINO,
    COL_JS_ESTADO, COL_JS_KGS, COL_JS_PRIORIDAD, COL_JS_EMPRESA,
]

# Columnas de la hoja GUIAS generada (y del detalle que muestra la app).
COL_GUIAS_GUIA = "Nro guía"
COL_GUIAS_FECHA = "Creación"
COL_GUIAS_CLIENTE = "Cliente"
COL_GUIAS_ORIGEN = "Origen"
COL_GUIAS_DESTINO = "Destino"
COL_GUIAS_ESTADO = "Estado"
COL_GUIAS_KG = "Kg"
COL_GUIAS_GRAVADO = "GRAVADO"
COL_GUIAS_IVA = "IVA"
COL_GUIAS_NO_GRAV = "NO GRAV"
COL_GUIAS_TOTAL = "TOTAL"


def load_jetsmart_export(path) -> pd.DataFrame:
    """Lee el export del sistema de guias y deja solo las filas de JetSmart.

    - Descarta columnas "Unnamed: N" (la planilla de trabajo de Anita trae
      tablas dinamicas pegadas a la derecha del export).
    - Descarta filas sin "# Guía" (fila de totales del pie).
    - Filtra por Empresa == JetSmart sin distinguir mayusculas (el export
      real trae "Jetsmart").
    """
    if hasattr(path, "seek"):
        path.seek(0)  # archivo subido en Streamlit: puede haberse leido antes
    df = pd.read_excel(path, sheet_name=0)
    faltantes = [c for c in _REQUIRED_COLUMNS if c not in df.columns]
    if faltantes:
        raise ValueError(
            "El archivo no parece ser el export de guías de JetSmart: faltan las columnas "
            + ", ".join(faltantes)
        )
    df = df[_REQUIRED_COLUMNS]
    df = df[df[COL_JS_GUIA].notna()]
    df = df[df[COL_JS_EMPRESA].astype(str).str.strip().str.upper() == "JETSMART"]
    return df.reset_index(drop=True)


def _creacion_to_timestamp(values: pd.Series) -> pd.Series:
    # El export trae "01/07/2026" como texto (dia primero); si algun dia
    # llega como fecha real de Excel, to_datetime la deja tal cual.
    return pd.to_datetime(values, dayfirst=True, errors="coerce")


def detect_jetsmart_period(df: pd.DataFrame) -> tuple[Period, pd.DataFrame]:
    """Periodo dominante por fecha de creacion, y las guias de otros meses.

    A diferencia de Avianca/Gol/LATAM, las guias de otros meses NO se
    excluyen: una guia creada el 30/6 que volo el 1/7 se liquida en julio
    (la regla es mes de vuelo, no de creacion) y el export no trae la fecha
    de vuelo. Se devuelven aparte solo para poder avisarlo en la interfaz.
    """
    fechas = _creacion_to_timestamp(df[COL_JS_CREACION])
    if fechas.isna().all():
        raise ValueError("No se pudo detectar el período: ninguna fila tiene una fecha de Creacion válida.")
    year_month = fechas.dt.to_period("M")
    dominante = year_month.value_counts().idxmax()
    period = (_MES_ABREV[dominante.month - 1], f"{dominante.year % 100:02d}")
    otros_meses = df[year_month != dominante]
    return period, otros_meses


def build_jetsmart_guias(df: pd.DataFrame) -> pd.DataFrame:
    """Hoja GUIAS: una fila por guia con Kg y el flete ($ Prioridad) gravado.

    "$ Prioridad" es el mismo valor que "Ingreso total flete" de la
    liquidacion a mano (verificado guia por guia en las 1925 guias en comun
    de julio 2026). Cga.Esp., Aduana y Mjo.Doc. no entran a esta
    liquidacion. IVA por guia = gravado x 21%, igual que la planilla.
    """
    prioridad = pd.to_numeric(df[COL_JS_PRIORIDAD], errors="coerce").fillna(0.0)
    iva = prioridad * JETSMART_IVA_RATE
    guias = pd.DataFrame({
        COL_GUIAS_GUIA: df[COL_JS_GUIA].map(lambda v: int(v) if isinstance(v, float) and v.is_integer() else v),
        COL_GUIAS_FECHA: df[COL_JS_CREACION],
        COL_GUIAS_CLIENTE: df[COL_JS_CLIENTE],
        COL_GUIAS_ORIGEN: df[COL_JS_ORIGEN].astype(str).str.strip().str.upper(),
        COL_GUIAS_DESTINO: df[COL_JS_DESTINO],
        COL_GUIAS_ESTADO: df[COL_JS_ESTADO],
        COL_GUIAS_KG: pd.to_numeric(df[COL_JS_KGS], errors="coerce").fillna(0.0),
        COL_GUIAS_GRAVADO: prioridad,
        COL_GUIAS_IVA: iva,
        COL_GUIAS_NO_GRAV: 0.0,
        COL_GUIAS_TOTAL: prioridad + iva,
    })
    return guias.reset_index(drop=True)


def tarifa_usd_kg(station: str) -> tuple[float, bool]:
    """(tarifa USD/kg, confirmada) para una estacion. Sin entrada en config
    usa el default, y nunca cuenta como confirmada."""
    tarifa = JETSMART_TARIFA_USD_KG_POR_ESTACION.get(station, JETSMART_TARIFA_USD_KG_DEFAULT)
    return tarifa, station in JETSMART_TARIFA_CONFIRMADA_STATIONS


def build_jetsmart_resumen(guias: pd.DataFrame, tipo_cambio: float, vuelos_inter: dict[str, int]) -> dict:
    """Calcula la hoja LIQUIDACION (y CVLP) a partir de la hoja GUIAS.

    tipo_cambio: TC del periodo (dato manual, ver app.py / --tc en main.py).
    vuelos_inter: {estacion: cantidad de vuelos internacionales} de los
    manifiestos, para "Commissions for sales - Inter" (dato manual).
    """
    if not tipo_cambio or tipo_cambio <= 0:
        raise ValueError("El tipo de cambio tiene que ser mayor a cero.")

    ventas_totales = float(guias[COL_GUIAS_TOTAL].sum())
    iva = -float(guias[COL_GUIAS_IVA].sum())
    ventas_netas = ventas_totales + iva
    comision_domestica = -ventas_netas * JETSMART_COMISION_DOMESTICA_RATE

    comisiones_inter = []
    for station, usd_por_vuelo in JETSMART_COMISION_INTER_USD_POR_VUELO.items():
        vuelos = int(vuelos_inter.get(station, 0) or 0)
        comisiones_inter.append({
            "estacion": station,
            "vuelos": vuelos,
            "usd_por_vuelo": usd_por_vuelo,
            "tc": tipo_cambio,
            "pesos": vuelos * usd_por_vuelo * tipo_cambio,
        })
    comision_inter = -sum(c["pesos"] for c in comisiones_inter)

    kg_por_estacion = guias.groupby(COL_GUIAS_ORIGEN)[COL_GUIAS_KG].sum()
    extra_stations = sorted(s for s in kg_por_estacion.index if s not in JETSMART_STATIONS)
    estaciones = []
    for station in JETSMART_STATIONS + extra_stations:
        kg = float(kg_por_estacion.get(station, 0.0))
        tarifa, confirmada = tarifa_usd_kg(station)
        estaciones.append({
            "estacion": station,
            "kg": kg,
            "tarifa_usd_kg": tarifa,
            "tc": tipo_cambio,
            "pesos": kg * tarifa * tipo_cambio,
            "tarifa_confirmada": confirmada,
        })
    gha_services = -sum(e["pesos"] for e in estaciones)

    total_wcs = (
        ventas_netas + comision_domestica + comision_inter + gha_services
        + JETSMART_IVA_SERVICIOS + JETSMART_IIBB
    )

    # Hoja CVLP (Cuenta de Venta y Liquido Producto): mismo dato, agrupado
    # en gravado / no gravado + IVA 21% sobre el saldo neto gravado.
    saldo_neto_gravado = ventas_netas + comision_domestica + gha_services
    saldo_neto_no_gravado = 0.0 + comision_inter
    total_neto = saldo_neto_gravado + saldo_neto_no_gravado
    iva_cvlp = saldo_neto_gravado * JETSMART_IVA_RATE

    return {
        "ventas_totales": ventas_totales,
        "iva": iva,
        "ventas_netas": ventas_netas,
        "comision_domestica": comision_domestica,
        "comision_inter": comision_inter,
        "gha_services": gha_services,
        "iva_servicios": JETSMART_IVA_SERVICIOS,
        "iibb": JETSMART_IIBB,
        "total_wcs": total_wcs,
        "estaciones": estaciones,
        "kg_total": sum(e["kg"] for e in estaciones),
        "comisiones_inter": comisiones_inter,
        "tipo_cambio": tipo_cambio,
        "cvlp": {
            "ventas_gravadas": ventas_netas,
            "saldo_neto_gravado": saldo_neto_gravado,
            "ventas_no_gravadas": 0.0,
            "saldo_neto_no_gravado": saldo_neto_no_gravado,
            "total_neto": total_neto,
            "iva_21": iva_cvlp,
            "total_final": total_neto + iva_cvlp,
        },
    }


def unconfirmed_tarifa_activity(resumen: dict) -> dict[str, float]:
    """{estacion: kg} de las estaciones con kilos y tarifa sin confirmar.
    Red de seguridad para el aviso amarillo de app.py (no bloquea nada)."""
    return {
        e["estacion"]: e["kg"]
        for e in resumen["estaciones"]
        if e["kg"] and not e["tarifa_confirmada"]
    }


def _period_last_day(period: Period) -> str:
    mes, anio = period
    month = _MES_ABREV.index(mes) + 1
    year = 2000 + int(anio)
    return f"{calendar.monthrange(year, month)[1]:02d}/{month:02d}/{year}"


_MONEY = "#,##0.00"


def _write_guias_sheet(writer: pd.ExcelWriter, guias: pd.DataFrame) -> None:
    guias.to_excel(writer, sheet_name="GUIAS", index=False)
    ws = writer.sheets["GUIAS"]
    total_row = len(guias) + 2
    bold = Font(bold=True)
    ws.cell(row=total_row, column=1, value="TOTAL").font = bold
    for col_idx, col in enumerate(guias.columns, start=1):
        if col in (COL_GUIAS_KG, COL_GUIAS_GRAVADO, COL_GUIAS_IVA, COL_GUIAS_NO_GRAV, COL_GUIAS_TOTAL):
            cell = ws.cell(row=total_row, column=col_idx, value=round(float(guias[col].sum()), 2))
            cell.font = bold
            cell.number_format = _MONEY


def _write_liquidacion_sheet(writer: pd.ExcelWriter, resumen: dict, period: Period) -> None:
    """Misma disposicion de celdas que la hoja LIQUIDACION a mano (A3:B13 y
    la tabla por estacion en E4:I16), para poder compararlas celda a celda."""
    ws = writer.book.create_sheet("LIQUIDACION")
    bold = Font(bold=True)

    def money(row, col, value, *, is_bold=False):
        cell = ws.cell(row=row, column=col, value=round(value, 2))
        cell.number_format = _MONEY
        if is_bold:
            cell.font = bold

    ws.cell(row=1, column=2, value=f"Beginning  Operation until {_period_last_day(period)}").font = bold

    left = [
        (3, "Ventas totales ", resumen["ventas_totales"]),
        (4, "IVA", resumen["iva"]),
        (6, "Ventas Netas ", resumen["ventas_netas"]),
        (7, "Commissions for sales - Domestic", resumen["comision_domestica"]),
        (8, "Commissions for sales - Inter", resumen["comision_inter"]),
        (9, "GHA Services", resumen["gha_services"]),
        (10, "IVA of Service And Comissions", resumen["iva_servicios"]),
        (11, "IIBB Tax (Over cost) (*)", resumen["iibb"]),
        (13, "Total collections to be delivered to WCS", resumen["total_wcs"]),
    ]
    for row, label, value in left:
        ws.cell(row=row, column=1, value=label)
        money(row, 2, value, is_bold=row in (6, 13))
    ws.cell(row=13, column=1).font = bold

    for col, header in zip(range(5, 11), ["Ex ", "Kg", "Rate", "TC", "Pesos", "Tarifa"]):
        ws.cell(row=4, column=col, value=header).font = bold
    row = 5
    for e in resumen["estaciones"]:
        ws.cell(row=row, column=5, value=e["estacion"])
        if e["kg"]:
            ws.cell(row=row, column=6, value=round(e["kg"], 3))
        ws.cell(row=row, column=7, value=e["tarifa_usd_kg"])
        ws.cell(row=row, column=8, value=e["tc"])
        money(row, 9, e["pesos"])
        ws.cell(row=row, column=10, value="confirmada" if e["tarifa_confirmada"] else "NO CONFIRMADA")
        row += 1
    ws.cell(row=row, column=5, value="Total").font = bold
    ws.cell(row=row, column=6, value=round(resumen["kg_total"], 3)).font = bold
    money(row, 9, -resumen["gha_services"], is_bold=True)

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["I"].width = 16
    ws.column_dimensions["J"].width = 16


def _write_cvlp_sheet(writer: pd.ExcelWriter, resumen: dict, period: Period) -> None:
    """Misma disposicion que la hoja CVLP a mano (B/C, filas 1-23)."""
    ws = writer.book.create_sheet("CVLP")
    bold = Font(bold=True)
    cvlp = resumen["cvlp"]
    ws.cell(row=1, column=2, value=f"LIQUIDACION {_period_last_day(period)[3:]}").font = bold
    rows = [
        (2, "VENTAS TOTALES", resumen["ventas_totales"]),
        (3, "IVA ", resumen["iva"]),
        (4, "VENTAS NETAS IVA", resumen["ventas_netas"]),
        (9, "VENTAS GRAVADAS", cvlp["ventas_gravadas"]),
        (10, "COMISION DE VENTAS DOMESTICO", resumen["comision_domestica"]),
        (11, "GHA SERVICIOS", resumen["gha_services"]),
        (12, "SALDO NETO GRAVADO", cvlp["saldo_neto_gravado"]),
        (14, "VENTAS NO GRAVADAS", cvlp["ventas_no_gravadas"]),
        (15, "RECUPERO GASTOS NO GRAVADOS", resumen["comision_inter"]),
        (16, "SALDO NETO NO GRAVADO", cvlp["saldo_neto_no_gravado"]),
        (19, "TOTAL NETO", cvlp["total_neto"]),
        (21, "IVA 21 S/ NETO GRAVADO", cvlp["iva_21"]),
        (23, "TOTAL FINAL", cvlp["total_final"]),
    ]
    for row, label, value in rows:
        ws.cell(row=row, column=2, value=label)
        cell = ws.cell(row=row, column=3, value=round(value, 2))
        cell.number_format = _MONEY
        if row in (12, 16, 19, 23):
            ws.cell(row=row, column=2).font = bold
            cell.font = bold
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 18


def _write_comisiones_inter_sheet(writer: pd.ExcelWriter, resumen: dict) -> None:
    """Reemplaza a las hojas "MANI <estacion>": solo el calculo, con la
    cantidad de vuelos cargada a mano (el detalle de manifiestos no viene
    en ningun export)."""
    ws = writer.book.create_sheet("COMISIONES INTER")
    bold = Font(bold=True)
    for col, header in enumerate(["Estación", "Vuelos", "USD por vuelo", "TC", "Pesos"], start=1):
        ws.cell(row=1, column=col, value=header).font = bold
    row = 2
    for c in resumen["comisiones_inter"]:
        ws.cell(row=row, column=1, value=c["estacion"])
        ws.cell(row=row, column=2, value=c["vuelos"])
        ws.cell(row=row, column=3, value=c["usd_por_vuelo"])
        ws.cell(row=row, column=4, value=c["tc"])
        cell = ws.cell(row=row, column=5, value=round(c["pesos"], 2))
        cell.number_format = _MONEY
        row += 1
    ws.cell(row=row, column=1, value="Total").font = bold
    cell = ws.cell(row=row, column=5, value=round(-resumen["comision_inter"], 2))
    cell.font = bold
    cell.number_format = _MONEY
    ws.cell(row=row + 2, column=1, value="Cantidad de vuelos cargada a mano a partir de los manifiestos del período.")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["E"].width = 16


def _write_cruce_sheets(writer: pd.ExcelWriter, excepciones: pd.DataFrame, incluidas_ultimo_dia: pd.DataFrame) -> None:
    """Resultado del cruce con Ariel (ver jetsmart_cruce.py): solo las
    excepciones, y aparte las guias incluidas por la regla de ultimo dia."""
    if excepciones.empty:
        pd.DataFrame({"Excepciones del cruce": ["Sin excepciones: todas las guías cruzaron limpio."]}).to_excel(
            writer, sheet_name="EXCEPCIONES CRUCE", index=False)
    else:
        excepciones.to_excel(writer, sheet_name="EXCEPCIONES CRUCE", index=False)
    if not incluidas_ultimo_dia.empty:
        incluidas_ultimo_dia.to_excel(writer, sheet_name="INCLUIDAS ULTIMO DIA", index=False)


def write_jetsmart_liquidacion(guias: pd.DataFrame, resumen: dict, period: Period, output_path, cruce: dict | None = None) -> None:
    """cruce: {"excepciones": DataFrame, "incluidas_ultimo_dia": DataFrame}
    del cruce con Ariel, si la liquidacion se armo a partir de el."""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_guias_sheet(writer, guias)
        _write_liquidacion_sheet(writer, resumen, period)
        _write_cvlp_sheet(writer, resumen, period)
        _write_comisiones_inter_sheet(writer, resumen)
        if cruce is not None:
            _write_cruce_sheets(writer, cruce["excepciones"], cruce["incluidas_ultimo_dia"])
