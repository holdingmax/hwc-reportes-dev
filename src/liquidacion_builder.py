"""Armado de la liquidacion formal de LATAM.

LATAM no encaja en el flujo simple de report_builder.py (una hoja por tipo
de cargo/estacion): necesita una liquidacion con 3 hojas -- Resumen
Facturacion, Detalle de Facturacion y Compensacion -- con una formula de
sub-facturas e IVA. Por eso vive en un modulo aparte, aunque reutiliza
utilidades genericas: _nonzero_mask/_fix_codigo_column de report_builder.py,
y deteccion/filtro de periodo de period_utils.py (compartida con el
reporte simple de Avianca/Gol).

Esta liquidacion esta validada unicamente para LATAM_LIQUIDACION_STATION
(EZE), contra el archivo de referencia real del cliente. No calcula
compensacion (ver nota en config.py junto a LATAM_SUBFACTURA_LA): esa parte
sigue siendo manual, tal como confirmo el cliente.
"""

import pandas as pd
from openpyxl.styles import Font

from src.config import (
    COL_AEROLINEA,
    COL_CLIENTE,
    COL_CODIGO,
    COL_COD_VUELO,
    COL_CONDICION,
    COL_ESTACION,
    COL_TIPO,
    COL_TPO_CAMBIO,
    LATAM_DETALLE_CHARGE_COLUMNS,
    LATAM_IVA_RATE,
    LATAM_LIQUIDACION_STATION,
    LATAM_SUBFACTURA_4M,
    LATAM_SUBFACTURA_LA,
)
from src.period_utils import Period, detect_period, excluded_by_period, filter_by_period
from src.report_builder import _fix_codigo_column, _nonzero_mask


def detect_latam_period_exclusions(df: pd.DataFrame) -> tuple[Period, pd.DataFrame]:
    """Periodo dominante de LATAM (todas las estaciones) y las filas que quedan afuera.

    Mismo criterio que detect_period_exclusions de report_builder.py para
    Avianca/Gol: se mira TODA la aerolinea, no solo LATAM_LIQUIDACION_STATION
    (EZE) -- una fila de COR/ROS/MDZ/NQN fuera de periodo tambien tiene que
    quedar reflejada en el aviso, aunque esa estacion ya este excluida de la
    liquidacion por otro motivo (ver detect_unconfirmed_station_activity).

    Expuesta para que el caller (ej. app.py) conozca el periodo detectado --
    para el nombre de archivo, para armar el aviso de filas excluidas, y
    para pasarselo explicitamente a build_latam_detalle/
    detect_unconfirmed_station_activity y evitar detectarlo mas de una vez.
    """
    latam_df = df[df[COL_AEROLINEA] == "LATAM"]
    period = detect_period(latam_df)
    excluded = excluded_by_period(latam_df, period)
    return period, excluded


def build_latam_detalle(df: pd.DataFrame, period: tuple[str, str] | None = None) -> pd.DataFrame:
    """Arma la hoja "Detalle de Facturacion" de LATAM para LATAM_LIQUIDACION_STATION.

    Una fila por guia con los datos identificatorios y las 4 columnas de
    cargo (nombres largos), copiadas tal cual del original sin ningun
    ajuste de tarifa -- LATAM usa el dato bruto, igual que Gol.

    period es (mes, anio) ej. ("JUL", "26"); si es None se autodetecta con
    detect_period().
    """
    latam_df = df[(df[COL_AEROLINEA] == "LATAM") & (df[COL_ESTACION] == LATAM_LIQUIDACION_STATION)]

    if period is None:
        period = detect_period(latam_df)
    latam_df = filter_by_period(latam_df, period)

    raw_columns = list(LATAM_DETALLE_CHARGE_COLUMNS.keys())
    mask = _nonzero_mask(latam_df, raw_columns)
    filtered = latam_df.loc[mask].copy()

    id_columns = [COL_CODIGO, COL_COD_VUELO, COL_CLIENTE, COL_CONDICION, COL_TIPO, COL_TPO_CAMBIO]
    detalle = filtered[id_columns + raw_columns].rename(columns=LATAM_DETALLE_CHARGE_COLUMNS)
    detalle = detalle.reset_index(drop=True)
    detalle = _fix_codigo_column(detalle)
    return detalle


def build_latam_resumen(detalle: pd.DataFrame) -> dict:
    """Calcula los totales del Resumen Facturacion a partir del Detalle.

    No aplica ningun ajuste de compensacion: usa los montos en bruto del
    Detalle (ver nota junto a LATAM_SUBFACTURA_LA en config.py).
    """
    sums = {
        col: float(pd.to_numeric(detalle[col], errors="coerce").fillna(0).sum())
        for col in LATAM_DETALLE_CHARGE_COLUMNS.values()
    }

    total_la = sum(sums[c] for c in LATAM_SUBFACTURA_LA)
    neto_gravado = sum(sums[c] for c in LATAM_SUBFACTURA_4M)
    iva = neto_gravado * LATAM_IVA_RATE
    total_4m = neto_gravado + iva
    total_periodo = total_la + total_4m

    return {
        "sums": sums,
        "total_la": total_la,
        "neto_gravado": neto_gravado,
        "iva": iva,
        "total_4m": total_4m,
        "total_periodo": total_periodo,
    }


def _write_resumen_sheet(writer: pd.ExcelWriter, resumen: dict) -> None:
    ws = writer.book.create_sheet("Resumen Facturación")

    bold = Font(bold=True)
    money_format = "#,##0.00"

    def write(row, label, value=None, *, is_bold=False):
        ws.cell(row=row, column=1, value=label)
        if is_bold:
            ws.cell(row=row, column=1).font = bold
        if value is not None:
            cell = ws.cell(row=row, column=3, value=round(value, 2))
            cell.number_format = money_format
            if is_bold:
                cell.font = bold

    sums = resumen["sums"]
    row = 1
    write(row, "LATAM AIRLINES", is_bold=True)
    row += 1
    for col in LATAM_SUBFACTURA_LA:
        write(row, col, sums[col])
        row += 1
    write(row, "TOTAL LA (sin IVA)", resumen["total_la"], is_bold=True)
    row += 2

    write(row, "LAN ARGENTINA", is_bold=True)
    row += 1
    for col in LATAM_SUBFACTURA_4M:
        write(row, col, sums[col])
        row += 1
    write(row, "Total IVA (21%, informativo)", resumen["iva"])
    row += 1
    write(row, "TOTAL 4M (con IVA)", resumen["total_4m"], is_bold=True)
    row += 2

    write(row, "TOTAL PERIODO", resumen["total_periodo"], is_bold=True)

    ws.column_dimensions["A"].width = 65
    ws.column_dimensions["C"].width = 18


def _write_compensacion_sheet(writer: pd.ExcelWriter) -> None:
    ws = writer.book.create_sheet("Compensación")
    ws.cell(row=1, column=1, value="Compensación").font = Font(bold=True, size=13)
    ws.cell(row=3, column=1, value="Cargar manualmente los casos de compensación del período.")
    ws.column_dimensions["A"].width = 65


def detect_unconfirmed_station_activity(
    df: pd.DataFrame,
    unconfirmed_stations: list[str],
    period: tuple[str, str] | None = None,
) -> dict[str, dict]:
    """Red de seguridad: detecta movimiento real de LATAM en estaciones sin validar.

    A diferencia de Avianca, esta liquidacion SI descarta esas filas (ver
    build_latam_detalle, que filtra a LATAM_LIQUIDACION_STATION): no entran
    a ningun total del Resumen. Esta funcion las detecta aparte para poder
    avisar que quedaron afuera con un monto real, en vez de omitirlas en
    silencio como hace hoy build_latam_detalle.

    Devuelve {estacion: {"filas": n, "total": suma de las 4 columnas de
    cargo}} solo para las estaciones con al menos una fila con cargo en el
    periodo detectado.

    OJO: se filtra por periodo antes de sumar por estacion, a proposito.
    Una fila fuera de periodo (de cualquier estacion) ya queda contemplada
    por detect_latam_period_exclusions -- si tambien se contara aca, el
    mismo monto aparecería duplicado en dos avisos distintos.
    """
    latam_df = df[df[COL_AEROLINEA] == "LATAM"]
    if period is None:
        period = detect_period(latam_df)
    latam_df = filter_by_period(latam_df, period)

    raw_columns = list(LATAM_DETALLE_CHARGE_COLUMNS.keys())
    activity: dict[str, dict] = {}
    for station in unconfirmed_stations:
        station_df = latam_df[latam_df[COL_ESTACION] == station]
        mask = _nonzero_mask(station_df, raw_columns)
        matched = station_df[mask]
        if matched.empty:
            continue
        total = float(sum(pd.to_numeric(matched[c], errors="coerce").fillna(0).sum() for c in raw_columns))
        activity[station] = {"filas": len(matched), "total": total}
    return activity


def write_liquidacion(detalle: pd.DataFrame, resumen: dict, output_path) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_resumen_sheet(writer, resumen)
        detalle.to_excel(writer, sheet_name="Detalle de Facturación", index=False)
        _write_compensacion_sheet(writer)
